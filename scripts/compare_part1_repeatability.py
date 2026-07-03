#!/usr/bin/env python3
import pandas as pd
import json
import numpy as np
import hashlib
from pathlib import Path
from typing import Dict, Any, List, Tuple
import openpyxl

def compute_sha256(file_path: Path) -> str:
    """Compute SHA-256 hash of a file."""
    sha256_hash = hashlib.sha256()
    with open(file_path, "rb") as f:
        for byte_block in iter(lambda: f.read(4096), b""):
            sha256_hash.update(byte_block)
    return sha256_hash.hexdigest()

def compare_csv_files(first_path: Path, second_path: Path, key_columns: List[str]) -> Dict[str, Any]:
    """Compare two CSV files with detailed analysis."""
    result = {
        "file": first_path.name,
        "exists_in_both": True,
        "first_sha256": "",
        "second_sha256": "",
        "sha256_match": False,
        "first_rows": 0,
        "second_rows": 0,
        "first_cols": 0,
        "second_cols": 0,
        "column_names_match": True,
        "column_order_match": True,
        "key_duplicates": {"first": False, "second": False},
        "missing_keys": [],
        "extra_keys": [],
        "text_columns_match": True,
        "text_differences": [],
        "numerical_analysis": {},
        "max_absolute_diff": 0.0,
        "mean_absolute_diff": 0.0,
        "diff_nonzero": 0,
        "diff_gt_1e12": 0,
        "diff_gt_1e9": 0,
        "null_mismatch_count": 0,
    }
    
    # Check file existence
    if not first_path.exists():
        result["exists_in_both"] = False
        result["error"] = "First file not found"
        return result
    
    if not second_path.exists():
        result["exists_in_both"] = False
        result["error"] = "Second file not found"
        return result
    
    # Compute SHA-256
    result["first_sha256"] = compute_sha256(first_path)
    result["second_sha256"] = compute_sha256(second_path)
    result["sha256_match"] = result["first_sha256"] == result["second_sha256"]
    
    # Read files
    try:
        first_df = pd.read_csv(first_path)
        second_df = pd.read_csv(second_path)
    except Exception as e:
        result["error"] = f"Error reading files: {str(e)}"
        return result
    
    result["first_rows"] = len(first_df)
    result["second_rows"] = len(second_df)
    result["first_cols"] = len(first_df.columns)
    result["second_cols"] = len(second_df.columns)
    
    # Column names and order
    first_cols = list(first_df.columns)
    second_cols = list(second_df.columns)
    
    if first_cols != second_cols:
        result["column_names_match"] = False
        result["column_names"] = {"first": first_cols, "second": second_cols}
    
    if first_cols == second_cols:
        result["column_order_match"] = True
    else:
        result["column_order_match"] = False
    
    # Key duplicates
    if key_columns:
        first_keys = first_df[key_columns].apply(tuple, axis=1)
        second_keys = second_df[key_columns].apply(tuple, axis=1)
        
        result["key_duplicates"]["first"] = first_keys.duplicated().any()
        result["key_duplicates"]["second"] = second_keys.duplicated().any()
        
        # Missing/extra keys
        first_key_set = set(first_keys)
        second_key_set = set(second_keys)
        
        result["missing_keys"] = list(first_key_set - second_key_set)
        result["extra_keys"] = list(second_key_set - first_key_set)
    
    # Merge on keys for comparison
    if key_columns:
        merged = first_df.merge(second_df, on=key_columns, suffixes=('_first', '_second'), how='outer', indicator=True)
    else:
        # If no keys, compare by index
        merged = first_df.merge(second_df, left_index=True, right_index=True, suffixes=('_first', '_second'), how='outer', indicator=True)
    
    # Text column comparison
    text_cols = first_df.select_dtypes(include=['object']).columns
    for col in text_cols:
        if f"{col}_first" in merged.columns and f"{col}_second" in merged.columns:
            first_col = merged[f"{col}_first"]
            second_col = merged[f"{col}_second"]
            
            # Check for null mismatches
            first_null = first_col.isna()
            second_null = second_col.isna()
            null_mismatch = (first_null != second_null)
            result["null_mismatch_count"] += null_mismatch.sum()
            
            if null_mismatch.any():
                result["text_columns_match"] = False
                for idx in merged[null_mismatch].index:
                    key_values = {k: merged.loc[idx, k] for k in key_columns} if key_columns else {}
                    result["text_differences"].append({
                        "column": col,
                        "index": idx,
                        "key_values": key_values,
                        "first_value": "NaN" if first_null.loc[idx] else str(first_col.loc[idx]),
                        "second_value": "NaN" if second_null.loc[idx] else str(second_col.loc[idx])
                    })
            
            # Compare non-null values
            mask = first_col.notna() & second_col.notna()
            if not (first_col[mask] == second_col[mask]).all():
                result["text_columns_match"] = False
                diff_mask = mask & (first_col != second_col)
                for idx in merged[diff_mask].index:
                    key_values = {k: merged.loc[idx, k] for k in key_columns} if key_columns else {}
                    result["text_differences"].append({
                        "column": col,
                        "index": idx,
                        "key_values": key_values,
                        "first_value": str(first_col.loc[idx]),
                        "second_value": str(second_col.loc[idx])
                    })
    
    # Numerical column comparison
    num_cols = first_df.select_dtypes(include=[np.number]).columns
    all_diffs = []
    
    for col in num_cols:
        if f"{col}_first" in merged.columns and f"{col}_second" in merged.columns:
            first_col = merged[f"{col}_first"]
            second_col = merged[f"{col}_second"]
            
            # Check for null mismatches
            first_null = first_col.isna()
            second_null = second_col.isna()
            null_mismatch = (first_null != second_null)
            result["null_mismatch_count"] += null_mismatch.sum()
            
            if null_mismatch.any():
                result["text_columns_match"] = False
                for idx in merged[null_mismatch].index:
                    key_values = {k: merged.loc[idx, k] for k in key_columns} if key_columns else {}
                    result["text_differences"].append({
                        "column": col,
                        "index": idx,
                        "key_values": key_values,
                        "first_value": "NaN" if first_null.loc[idx] else str(first_col.loc[idx]),
                        "second_value": "NaN" if second_null.loc[idx] else str(second_col.loc[idx])
                    })
            
            # Handle NaN - compare only where both are non-null
            mask = first_col.notna() & second_col.notna()
            if mask.sum() > 0:
                diff = np.abs(first_col[mask] - second_col[mask])
                all_diffs.extend(diff.tolist())
                
                col_max_diff = diff.max() if len(diff) > 0 else 0.0
                col_mean_diff = diff.mean() if len(diff) > 0 else 0.0
                
                result["numerical_analysis"][col] = {
                    "max_diff": float(col_max_diff),
                    "mean_diff": float(col_mean_diff),
                    "diff_nonzero": int((diff != 0).sum()),
                    "diff_gt_1e12": int((diff > 1e-12).sum()),
                    "diff_gt_1e9": int((diff > 1e-9).sum())
                }
    
    if all_diffs:
        result["max_absolute_diff"] = float(max(all_diffs))
        result["mean_absolute_diff"] = float(np.mean(all_diffs))
        result["diff_nonzero"] = sum(1 for d in all_diffs if d != 0)
        result["diff_gt_1e12"] = sum(1 for d in all_diffs if d > 1e-12)
        result["diff_gt_1e9"] = sum(1 for d in all_diffs if d > 1e-9)
    
    return result

def compare_json_files(first_path: Path, second_path: Path) -> Dict[str, Any]:
    """Compare two JSON files field by field."""
    result = {
        "file": first_path.name,
        "exists_in_both": True,
        "first_sha256": "",
        "second_sha256": "",
        "sha256_match": False,
        "fields_match": True,
        "differences": []
    }
    
    if not first_path.exists():
        result["exists_in_both"] = False
        result["error"] = "First file not found"
        return result
    
    if not second_path.exists():
        result["exists_in_both"] = False
        result["error"] = "Second file not found"
        return result
    
    # Compute SHA-256
    result["first_sha256"] = compute_sha256(first_path)
    result["second_sha256"] = compute_sha256(second_path)
    result["sha256_match"] = result["first_sha256"] == result["second_sha256"]
    
    try:
        with open(first_path, 'r') as f:
            first_data = json.load(f)
        with open(second_path, 'r') as f:
            second_data = json.load(f)
    except Exception as e:
        result["error"] = f"Error reading files: {str(e)}"
        return result
    
    # Compare field by field
    all_keys = set(first_data.keys()) | set(second_data.keys())
    
    for key in all_keys:
        first_val = first_data.get(key)
        second_val = second_data.get(key)
        
        if first_val != second_val:
            result["fields_match"] = False
            result["differences"].append({
                "field": key,
                "first_value": str(first_val),
                "second_value": str(second_val)
            })
    
    return result

def compare_excel_files(first_path: Path, second_path: Path) -> Dict[str, Any]:
    """Compare two Excel files."""
    result = {
        "file": first_path.name,
        "exists_in_both": True,
        "first_sha256": "",
        "second_sha256": "",
        "sha256_match": False,
        "sheet_names_match": True,
        "sheet_order_match": True,
        "sheet_details": {},
        "max_absolute_diff": 0.0,
        "text_differences": [],
        "null_mismatch_count": 0,
    }
    
    if not first_path.exists():
        result["exists_in_both"] = False
        result["error"] = "First file not found"
        return result
    
    if not second_path.exists():
        result["exists_in_both"] = False
        result["error"] = "Second file not found"
        return result
    
    # Compute SHA-256
    result["first_sha256"] = compute_sha256(first_path)
    result["second_sha256"] = compute_sha256(second_path)
    result["sha256_match"] = result["first_sha256"] == result["second_sha256"]
    
    try:
        first_wb = openpyxl.load_workbook(first_path, data_only=True)
        second_wb = openpyxl.load_workbook(second_path, data_only=True)
    except Exception as e:
        result["error"] = f"Error reading files: {str(e)}"
        return result
    
    first_sheets = first_wb.sheetnames
    second_sheets = second_wb.sheetnames
    
    if first_sheets != second_sheets:
        result["sheet_names_match"] = False
        result["sheet_order_match"] = False
    else:
        result["sheet_order_match"] = True
    
    all_sheets = set(first_sheets) | set(second_sheets)
    
    for sheet_name in all_sheets:
        sheet_result = {
            "exists_in_both": True,
            "first_rows": 0,
            "second_rows": 0,
            "first_cols": 0,
            "second_cols": 0,
            "column_names_match": True,
            "max_diff": 0.0,
            "text_differences": [],
            "null_mismatches": 0
        }
        
        if sheet_name not in first_sheets:
            sheet_result["exists_in_both"] = False
            result["sheet_details"][sheet_name] = sheet_result
            continue
        
        if sheet_name not in second_sheets:
            sheet_result["exists_in_both"] = False
            result["sheet_details"][sheet_name] = sheet_result
            continue
        
        first_ws = first_wb[sheet_name]
        second_ws = second_wb[sheet_name]
        
        sheet_result["first_rows"] = first_ws.max_row
        sheet_result["second_rows"] = second_ws.max_row
        sheet_result["first_cols"] = first_ws.max_column
        sheet_result["second_cols"] = second_ws.max_column
        
        # Read data for comparison
        first_data = []
        for row in first_ws.iter_rows(values_only=True):
            first_data.append(row)
        
        second_data = []
        for row in second_ws.iter_rows(values_only=True):
            second_data.append(row)
        
        # Compare cell by cell
        max_rows = max(len(first_data), len(second_data))
        max_cols = max(len(first_data[0]) if first_data else 0, len(second_data[0]) if second_data else 0)
        
        all_diffs = []
        
        for row_idx in range(max_rows):
            for col_idx in range(max_cols):
                first_val = first_data[row_idx][col_idx] if row_idx < len(first_data) and col_idx < len(first_data[row_idx]) else None
                second_val = second_data[row_idx][col_idx] if row_idx < len(second_data) and col_idx < len(second_data[row_idx]) else None
                
                # Check for null mismatches
                first_null = pd.isna(first_val)
                second_null = pd.isna(second_val)
                
                if first_null != second_null:
                    sheet_result["null_mismatches"] += 1
                    result["null_mismatch_count"] += 1
                    sheet_result["text_differences"].append({
                        "row": row_idx + 1,
                        "col": col_idx + 1,
                        "first_value": "NaN" if first_null else str(first_val),
                        "second_value": "NaN" if second_null else str(second_val)
                    })
                elif not first_null and not second_null:
                    # Both non-null, compare values
                    if isinstance(first_val, (int, float)) and isinstance(second_val, (int, float)):
                        diff = abs(first_val - second_val)
                        all_diffs.append(diff)
                        if diff > 0:
                            sheet_result["max_diff"] = max(sheet_result["max_diff"], diff)
                    elif first_val != second_val:
                        sheet_result["text_differences"].append({
                            "row": row_idx + 1,
                            "col": col_idx + 1,
                            "first_value": str(first_val),
                            "second_value": str(second_val)
                        })
        
        sheet_result["max_diff"] = float(sheet_result["max_diff"])
        result["max_absolute_diff"] = max(result["max_absolute_diff"], sheet_result["max_diff"])
        result["sheet_details"][sheet_name] = sheet_result
    
    return result

def main():
    base_dir = Path(__file__).resolve().parents[1]
    first_dir = base_dir / "results/part1_full_reproduction"
    second_dir = base_dir / "results/part1_full_reproduction_repeat"
    
    # Define files to compare with their key columns
    files_to_compare = {
        "dataset_profile.csv": ["project"],
        "repeated_all_results.csv": ["experiment", "target_project", "seed", "model"],
        "validation_log.csv": ["experiment", "target_project", "seed", "candidate", "mode"],
        "repeated_summary_mean_std.csv": ["experiment", "model"],
    }
    
    results = {}
    
    # Compare CSV files
    for filename, key_cols in files_to_compare.items():
        first_path = first_dir / filename
        second_path = second_dir / filename
        results[filename] = compare_csv_files(first_path, second_path, key_cols)
    
    # Compare JSON file
    json_result = compare_json_files(first_dir / "decision_metadata.json", second_dir / "decision_metadata.json")
    results["decision_metadata.json"] = json_result
    
    # Compare Excel file
    excel_result = compare_excel_files(first_dir / "repeated_results_workbook.xlsx", second_dir / "repeated_results_workbook.xlsx")
    results["repeated_results_workbook.xlsx"] = excel_result
    
    # Determine overall status
    has_text_diffs = any(
        r.get("text_columns_match") == False or 
        r.get("fields_match") == False or
        len(r.get("text_differences", [])) > 0 or
        len(r.get("differences", [])) > 0 or
        len(r.get("missing_keys", [])) > 0 or
        len(r.get("extra_keys", [])) > 0 or
        r.get("null_mismatch_count", 0) > 0 or
        not r.get("sheet_names_match", True) or
        not r.get("sheet_order_match", True)
        for r in results.values()
    )
    
    has_numerical_diffs = any(r.get("max_absolute_diff", 0.0) > 0 for r in results.values())
    max_diff = max(r.get("max_absolute_diff", 0.0) for r in results.values())
    
    if has_text_diffs:
        overall_status = "materially_different"
    elif has_numerical_diffs and max_diff > 1e-9:
        overall_status = "materially_different"
    elif has_numerical_diffs and max_diff > 1e-12:
        overall_status = "numerically_equivalent"
    else:
        overall_status = "exact_match"
    
    results["overall_status"] = overall_status
    
    # Convert numpy types to native Python types for JSON serialization
    def convert_to_native(obj):
        if isinstance(obj, dict):
            return {k: convert_to_native(v) for k, v in obj.items()}
        elif isinstance(obj, list):
            return [convert_to_native(item) for item in obj]
        elif isinstance(obj, (np.integer, np.int64, np.int32)):
            return int(obj)
        elif isinstance(obj, (np.floating, np.float64, np.float32)):
            return float(obj)
        elif isinstance(obj, np.bool_):
            return bool(obj)
        elif isinstance(obj, np.ndarray):
            return convert_to_native(obj.tolist())
        else:
            return obj
    
    results_native = convert_to_native(results)
    
    # Save JSON report
    json_report_path = base_dir / "reports/part1_repeatability_comparison.json"
    with open(json_report_path, 'w') as f:
        json.dump(results_native, f, indent=2)
    
    # Save Markdown report
    md_report_path = base_dir / "reports/part1_repeatability_comparison.md"
    with open(md_report_path, 'w') as f:
        f.write("# Part 1 Section 5H: Repeatability Comparison Report\n\n")
        f.write(f"## Overall Status\n\n")
        f.write(f"**{overall_status}**\n\n")
        f.write(f"## Maximum Numerical Difference\n\n")
        f.write(f"{max_diff:.2e}\n\n")
        
        for filename, result in results.items():
            if filename == "overall_status":
                continue
            
            f.write(f"## {filename}\n\n")
            
            if not result.get("exists_in_both"):
                f.write(f"**Error**: {result.get('error', 'File not found')}\n\n")
                continue
            
            if "sha256_match" in result:
                f.write(f"- **SHA-256 match**: {result['sha256_match']}\n")
                f.write(f"- **First SHA-256**: {result['first_sha256']}\n")
                f.write(f"- **Second SHA-256**: {result['second_sha256']}\n")
            
            f.write(f"- **First rows**: {result.get('first_rows', 'N/A')}\n")
            f.write(f"- **Second rows**: {result.get('second_rows', 'N/A')}\n")
            f.write(f"- **First columns**: {result.get('first_cols', 'N/A')}\n")
            f.write(f"- **Second columns**: {result.get('second_cols', 'N/A')}\n")
            f.write(f"- **Column names match**: {result.get('column_names_match', 'N/A')}\n")
            f.write(f"- **Column order match**: {result.get('column_order_match', 'N/A')}\n")
            
            if "key_duplicates" in result:
                f.write(f"- **Key duplicates (first)**: {result['key_duplicates']['first']}\n")
                f.write(f"- **Key duplicates (second)**: {result['key_duplicates']['second']}\n")
                f.write(f"- **Missing keys**: {len(result['missing_keys'])}\n")
                f.write(f"- **Extra keys**: {len(result['extra_keys'])}\n")
            
            if "text_columns_match" in result:
                f.write(f"- **Text columns match**: {result['text_columns_match']}\n")
                if result['text_differences']:
                    f.write(f"- **Text differences**: {len(result['text_differences'])}\n")
            
            if "fields_match" in result:
                f.write(f"- **Fields match**: {result['fields_match']}\n")
                if result['differences']:
                    f.write(f"- **Field differences**: {len(result['differences'])}\n")
            
            if "sheet_names_match" in result:
                f.write(f"- **Sheet names match**: {result['sheet_names_match']}\n")
                f.write(f"- **Sheet order match**: {result['sheet_order_match']}\n")
            
            if "max_absolute_diff" in result:
                f.write(f"- **Max absolute difference**: {result['max_absolute_diff']:.2e}\n")
                f.write(f"- **Mean absolute difference**: {result.get('mean_absolute_diff', 0.0):.2e}\n")
                f.write(f"- **Non-zero differences**: {result.get('diff_nonzero', 0)}\n")
                f.write(f"- **Differences > 1e-12**: {result.get('diff_gt_1e12', 0)}\n")
                f.write(f"- **Differences > 1e-9**: {result.get('diff_gt_1e9', 0)}\n")
            
            if "null_mismatch_count" in result:
                f.write(f"- **Null mismatches**: {result['null_mismatch_count']}\n")
            
            f.write("\n")
    
    print(f"Comparison complete. Overall status: {overall_status}")
    print(f"JSON report saved to: {json_report_path}")
    print(f"Markdown report saved to: {md_report_path}")

if __name__ == "__main__":
    main()
